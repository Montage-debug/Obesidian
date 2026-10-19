from pathlib import Path
from zipfile import ZipFile
import hashlib,json,re
from docx import Document
from lxml import etree
from pypdf import PdfReader

B=Path(__file__).parent
manifest=json.loads((B/'build_manifest.json').read_text(encoding='utf-8'))
out=Path(manifest['output']);d=Document(out)
pdf=PdfReader(B/'draft_final.pdf')
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
z=ZipFile(out);root=etree.fromstring(z.read('word/document.xml'))
checks={}
checks['pages']=len(pdf.pages)
checks['tables']=len(d.tables)
checks['images']=len(d.inline_shapes)
checks['editable_OMML_equations']=len(root.findall('.//m:oMath',ns))
assert checks['pages']==68 and checks['tables']==19 and checks['images']==15 and checks['editable_OMML_equations']==48
headings=[p.text for p in d.paragraphs if p.style.name=='Heading 1']
chapters=[t for t in headings if re.match(r'第[1-6]章',t.replace(' ',''))]
assert len(chapters)==6,chapters
checks['six_chapters']=chapters
fields=' '.join(e.text or '' for e in root.findall('.//w:instrText',ns))
assert 'TOC ' in fields
checks['automatic_toc_field']=True
toc_text=''.join(pdf.pages[i].extract_text() for i in (3,4,5))
assert '第1章' in toc_text.replace(' ','')
assert not re.search(r'目录[.\s…]+III',toc_text)
checks['toc_excludes_cover_and_itself']=True
gap_ids=[m['id'] for m in manifest['missing']]
assert len(gap_ids)==len(set(gap_ids))==30
texts=[p.text for p in d.paragraphs]
for gid in gap_ids:
    p=next(p for p in d.paragraphs if re.match(r'【待[补核]'+gid+'：',p.text))
    assert all(str(r.font.color.rgb)=='C00000' for r in p.runs if r.text)
index=d.tables[-1]
assert len(index.rows)==31
assert all(any(gid in row.cells[0].text for row in index.rows) for gid in gap_ids)
checks['red_placeholders_and_index']=gap_ids
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if '____' in r.text:assert str(r.font.color.rgb)=='C00000'
checks['blank_table_values_red']=True
refs=[p.text for p in d.paragraphs if re.match(r'^\[\d+\]',p.text)]
assert [int(re.match(r'^\[(\d+)\]',t)[1]) for t in refs]==list(range(1,18))
body=''.join((B/f'chapter{i}.txt').read_text(encoding='utf-8') for i in range(1,7))
cited=set(int(v) for v in re.findall(r'\[(\d+)\]',body))
for a,b in re.findall(r'\[(\d+)-(\d+)\]',body):cited.update(range(int(a),int(b)+1))
assert cited==set(range(1,18)),cited
checks['references_numbered_and_cited']=17
pdftext='\n'.join(p.extract_text() or '' for p in pdf.pages)
for bad in ['未定义书签','未找到引用源','Error! Reference','@eq|','@fig|','@table|','@gap|','turn0search']:
    assert bad not in pdftext,bad
checks['no_broken_fields_or_authoring_markers']=True
originals=json.loads((B/'source_manifest.json').read_text(encoding='utf-8'))
verified=[]
for name,item in originals.items():
    if isinstance(item,dict) and 'path' in item and 'sha256' in item:
        f=Path(item['path'])
        assert hashlib.sha256(f.read_bytes()).hexdigest()==item['sha256'],str(f)
        verified.append(str(f))
checks['original_documents_unchanged']=verified
changes={}
for prior,current in [('render2','render3'),('render3','render_final')]:
    changes[prior+' to '+current]=[i for i in range(1,69) if (B/prior/f'page-{i}.png').read_bytes()!=(B/current/f'page-{i}.png').read_bytes()]
checks['render_changes']=changes
checks['visual_review']={'pages_1_to_16':'render2 reviewed; changed TOC pages 4-6 reviewed in render3','pages_17_to_68':'all reviewed in render3','final_changed_pages':[46,57,58,59,61],'all_final_pages_covered':True}
checks['output_sha256']=hashlib.sha256(out.read_bytes()).hexdigest()
checks['output_bytes']=out.stat().st_size
(B/'qa_final.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'PASS':True,'pages':checks['pages'],'red_items':len(gap_ids),'equations':checks['editable_OMML_equations'],'references':17,'unchanged_sources':len(verified),'bytes':checks['output_bytes']}))
